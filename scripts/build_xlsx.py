"""Build the current-state spreadsheet for the custom set.

Usage: python scripts/build_xlsx.py [out.xlsx]      (default: out/set.xlsx)

Reads design/skeleton.yaml, cards/*.yaml, baseline/original-pool.json and
design/open-questions.yaml, and writes one workbook:
  Set             every skeleton slot with the card (if any) filling it
  Summary         progress, distribution vs the original pool, and rule checks
  Baseline        the original 133-card pool for reference
  Open Questions  unresolved design decisions
"""
import collections
import json
import sys
from pathlib import Path

import yaml
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

ROOT = Path(__file__).resolve().parent.parent
COLORS = ["white", "blue", "black", "red", "green", "colorless"]
RARITIES = ["common", "uncommon", "rare", "mythic"]
FILL = {
    "white": "F5F0DC", "blue": "CFE2F3", "black": "BFBFBF", "red": "F4CCCC", "green": "D9EAD3", "colorless": "EDEDED",
}
HEADER_FILL = PatternFill("solid", fgColor="333333")
HEADER_FONT = Font(bold=True, color="FFFFFF")


def load_yaml(path):
    with open(path, encoding="utf8") as f:
        return yaml.safe_load(f)


def timing_str(c):
    parts = [n for k, n in (("to_play", "to-play"), ("while_in_play", "while"), ("after_playing", "after")) if c.get(k)]
    return "+".join(parts) if parts else "none"


def value_str(c):
    if c.get("alt_value") is None:
        return str(c.get("base_value", ""))
    return "%s->%s" % (c.get("base_value"), c["alt_value"])


def load_cards():
    cards = []
    for p in sorted((ROOT / "cards").glob("*.yaml")):
        c = load_yaml(p)
        if not c:
            continue
        c["_file"] = p.name
        cards.append(c)
    return cards


def write_table(ws, header, rows, widths=None, color_col=None):
    ws.append(header)
    for cell in ws[1]:
        cell.fill, cell.font = HEADER_FILL, HEADER_FONT
        cell.alignment = Alignment(vertical="center", wrap_text=True)
    for r in rows:
        ws.append(r)
    ws.freeze_panes = "A2"
    for i, w in enumerate(widths or [], 1):
        ws.column_dimensions[get_column_letter(i)].width = w
    for row in ws.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = Alignment(vertical="top", wrap_text=True)
        if color_col is not None:
            fill = FILL.get(str(row[color_col].value))
            if fill:
                row[color_col].fill = PatternFill("solid", fgColor=fill)
    ws.auto_filter.ref = ws.dimensions


def main():
    out = Path(sys.argv[1]) if len(sys.argv) > 1 else ROOT / "out" / "set.xlsx"
    out.parent.mkdir(parents=True, exist_ok=True)

    skeleton = load_yaml(ROOT / "design" / "skeleton.yaml")
    slots = skeleton["slots"]
    meta = skeleton["set"]
    cards = load_cards()
    baseline = json.loads((ROOT / "baseline" / "original-pool.json").read_text(encoding="utf8"))
    base_names = {c["name"] for c in baseline}
    questions = (load_yaml(ROOT / "design" / "open-questions.yaml") or {}).get("questions", [])

    slot_by_id = {s["id"]: s for s in slots}
    by_slot = collections.defaultdict(list)
    unslotted = []
    problems = []
    for c in cards:
        sid = c.get("slot")
        if not sid:
            unslotted.append(c)
        elif sid not in slot_by_id:
            problems.append("%s: unknown slot %s" % (c["_file"], sid))
            unslotted.append(c)
        else:
            by_slot[sid].append(c)
            s = slot_by_id[sid]
            for k in ("color", "rarity"):
                if c.get(k) != s[k]:
                    problems.append("%s: %s is %s but slot %s expects %s" % (c["_file"], k, c.get(k), sid, s[k]))
        if c.get("reprint_of") and c["reprint_of"] not in base_names:
            problems.append("%s: reprint_of %r is not in the original pool" % (c["_file"], c["reprint_of"]))
    for sid, cs in by_slot.items():
        if len(cs) > 1:
            problems.append("slot %s claimed by %d cards: %s" % (sid, len(cs), ", ".join(c["_file"] for c in cs)))
    names = collections.Counter(c.get("name") for c in cards)
    for n, k in names.items():
        if k > 1:
            problems.append("duplicate card name %r" % n)
    reprints = [c for c in cards if c.get("reprint_of")]
    if len(reprints) > meta["reprint_cap"]:
        problems.append("reprints: %d exceeds cap of %d (%.0f%% of %d)" % (
            len(reprints), meta["reprint_cap"], 100 * meta["max_reprint_fraction"], meta["target_size"]))

    wb = Workbook()

    # --- Set sheet
    ws = wb.active
    ws.title = "Set"
    rows = []
    for s in slots:
        filled = by_slot.get(s["id"], [])
        c = filled[0] if filled else {}
        rows.append([
            s["id"], s["color"], s["rarity"], s["role"], s["value"], s["timing"], s.get("reprint_candidate", ""),
            c.get("status", "open") if c else "open", c.get("name", ""),
            value_str(c) if c else "", timing_str(c) if c else "", c.get("rules_text", ""),
            c.get("reprint_of", "") or "", c.get("impl_tier", ""), c.get("designer", ""), c.get("notes", "") if c else s["notes"],
        ])
    for c in unslotted:
        rows.append([
            "(none)", c.get("color", ""), c.get("rarity", ""), "", "", "", "", c.get("status", ""), c.get("name", ""),
            value_str(c), timing_str(c), c.get("rules_text", ""), c.get("reprint_of", "") or "",
            c.get("impl_tier", ""), c.get("designer", ""), c.get("notes", ""),
        ])
    write_table(ws, ["Slot", "Color", "Rarity", "Role", "Target value", "Target timing", "Reprint candidate", "Status", "Name",
                     "Value", "Timing", "Rules text", "Reprint of", "Impl tier", "Designer", "Notes"],
                rows, [8, 10, 10, 22, 11, 12, 14, 10, 18, 9, 12, 60, 14, 9, 10, 50], color_col=1)

    # --- Summary sheet
    sm = wb.create_sheet("Summary")
    sm.append(["Set:", meta.get("name", "TBD")])
    candidates = sum(1 for s in slots if s.get("reprint_candidate"))
    sm.append(["Slots", len(slots), "Cards", len(cards), "Reprints", len(reprints), "Reprint cap", meta["reprint_cap"],
               "Reprint candidates", candidates])
    sm.append([])
    sm.append(["Progress (filled / slots)"] + RARITIES + ["total"])
    for col in COLORS:
        row = [col]
        tf = ts = 0
        for r in RARITIES:
            ns = sum(1 for s in slots if s["color"] == col and s["rarity"] == r)
            nf = sum(1 for s in slots if s["color"] == col and s["rarity"] == r and by_slot.get(s["id"]))
            tf, ts = tf + nf, ts + ns
            row.append("%d / %d" % (nf, ns))
        row.append("%d / %d" % (tf, ts))
        sm.append(row)
    sm.append([])
    sm.append(["Status counts"] + sorted({c.get("status", "?") for c in cards}))
    stat = collections.Counter(c.get("status", "?") for c in cards)
    sm.append(["cards"] + [stat[k] for k in sorted(stat)])
    sm.append([])
    sm.append(["Average printed value (filled cards vs original pool)", "this set", "original"])
    for r in RARITIES:
        mine = [c["base_value"] for c in cards if c.get("rarity") == r and c.get("base_value") is not None]
        orig = [c["base_value"] for c in baseline if c["rarity"] == r]
        sm.append([r, round(sum(mine) / len(mine), 2) if mine else "", round(sum(orig) / len(orig), 2)])
    sm.append([])
    sm.append(["Ability timing (cards)", "this set", "original"])
    for k, label in (("to_play", "to-play"), ("while_in_play", "while in play"), ("after_playing", "after playing")):
        sm.append([label, sum(1 for c in cards if c.get(k)), sum(1 for c in baseline if c[k])])
    sm.append(["none (vanilla)", sum(1 for c in cards if not (c.get("to_play") or c.get("while_in_play") or c.get("after_playing"))),
               sum(1 for c in baseline if not (c["to_play"] or c["while_in_play"] or c["after_playing"]))])
    sm.append([])
    sm.append(["Rule checks"])
    for p in problems or ["none: all checks pass"]:
        sm.append([p])
    sm.column_dimensions["A"].width = 52
    for col in "BCDEFGH":
        sm.column_dimensions[col].width = 14
    for row in sm.iter_rows():
        if row and isinstance(row[0].value, str) and ("(" in row[0].value and ")" in row[0].value or row[0].value in ("Rule checks",)):
            row[0].font = Font(bold=True)

    # --- Baseline sheet
    bs = wb.create_sheet("Baseline")
    write_table(bs, ["#", "Name", "Color", "Rarity", "Value", "Timing", "Rules text"],
                [[c["id"], c["name"], c["color"], c["rarity"], value_str(c), timing_str(c), c["rules_text"]] for c in baseline],
                [5, 18, 10, 10, 9, 14, 90], color_col=2)

    # --- Open questions sheet
    oq = wb.create_sheet("Open Questions")
    write_table(oq, ["Id", "Status", "Question", "Notes / options"],
                [[q.get("id", ""), q.get("status", "open"), q.get("question", ""), q.get("notes", "")] for q in questions],
                [8, 10, 60, 90])

    wb.save(out)
    print("wrote %s: %d slots, %d cards, %d problem(s)" % (out, len(slots), len(cards), len(problems)))
    for p in problems:
        print("  PROBLEM:", p)


if __name__ == "__main__":
    main()
